# -*- coding: utf-8 -*-
"""生成 pure_pursuit.cpp 状态机说明 Word 文档
正文：中文宋体小四(12pt)，英文 Times New Roman 小四(12pt)
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CN_FONT = "宋体"
EN_FONT = "Times New Roman"
SIZE = Pt(12)  # 小四


def set_run_font(run, bold=False):
    run.font.name = EN_FONT
    run.font.size = SIZE
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), EN_FONT)
    rFonts.set(qn("w:hAnsi"), EN_FONT)
    rFonts.set(qn("w:eastAsia"), CN_FONT)


def add_para(doc, text, bold=False, align=None, indent=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent and align is None:
        p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进2字符
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, bold=True)
    if level == 1:
        run.font.color.rgb = RGBColor(0x1F, 0x3F, 0x7F)
    return p


def set_cell(cell, text, bold=False, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, bold)
    if shade:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), shade)
        cell._element.get_or_add_tcPr().append(shd)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, bold=True, shade="D9E2F3")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            set_cell(t.rows[i].cells[j], v)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = w
    return t


doc = Document()

# 全局默认样式也设置字体，防止空段落串字体
style = doc.styles["Normal"]
style.font.name = EN_FONT
style.font.size = SIZE
style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

# ===================== 标题 =====================
add_para(doc, "pure_pursuit.cpp 状态机及进入状态机条件说明",
         bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
add_para(doc, "（基于 crawler_control/src/pure_pursuit/src/pure_pursuit.cpp）",
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

# ===================== 一、概述 =====================
add_heading(doc, "一、状态机概述")
add_para(doc, "pure_pursuit 节点是纯跟踪（Pure Pursuit）路径跟踪控制节点，其核心是位于 "
         "state_machine_run() 函数中的运行状态机。节点启动后在构造函数中直接调用 "
         "state_machine_run()，以 20 Hz 的频率循环执行。")
add_para(doc, "状态机的状态由枚举 RunStateValue 定义（见 pure_pursuit.h 第43~48行），"
         "共 4 个状态，存放于原子变量 running_state_ 中：")
add_table(doc,
          ["状态", "枚举值", "含义"],
          [["Stop", "0", "停车：发布零速度、零角速度指令"],
           ["Follow", "1", "跟线：执行纯跟踪算法跟踪路径（节点启动后的初始状态）"],
           ["Turn", "2", "原地转弯：在转弯点（gear=2）处原地旋转对准下一段航向"],
           ["Reverse", "3", "倒车：紧急避障时以 -0.15 m/s 匀速倒车"]])
add_para(doc, "每个循环周期依次执行：① 接收回调（spinOnce）；② 前置安全检查（数据超时、避障预警），"
         "满足条件时直接跳转到状态机入口（goto STATEMACHINE）；③ 若处于 Turn 状态，刷新预瞄点角度；"
         "④ 查找最近路径点并计算预瞄点；⑤ 进入 switch 状态机，按当前状态计算并发布控制指令。")

# ===================== 二、各状态行为 =====================
add_heading(doc, "二、各状态行为说明")
add_heading(doc, "1. Stop（停车状态）", level=2)
add_para(doc, "清零 follow_count_、turn_count_ 计数器，发布线速度 0、角速度 0 的停车指令，"
         "随后立即将状态切回 Follow。因此 Stop 是一个“停一拍”的瞬时状态：每循环一次只停一个周期，"
         "若进入条件仍然成立，下一周期会被再次置为 Stop，从而实现持续停车。")
add_heading(doc, "2. Follow（跟线状态）", level=2)
add_para(doc, "正常运行状态。计算预瞄距离，调用 calculate_PurePursuit() 按纯跟踪算法计算线速度、"
         "角速度并发布。内含卡死检测：若预瞄点不变且预瞄距离一个周期内缩小不足 0.05 m，"
         "follow_count_ 累加；超过 300 次（约 15 s）判定小车卡住，强制以 1.0 m/s 前进 20 个周期脱困。")
add_para(doc, "若已到达预瞄点（预瞄距离 < 0.05 m 或预瞄点落到车身后方 local_x < 0.03）"
         "且该预瞄点为转弯点（gear=2），则停车并切换到 Turn 状态。")
add_heading(doc, "3. Turn（原地转弯状态）", level=2)
add_para(doc, "小车在转弯点处原地旋转，线速度恒为 0，角速度由 calculate_rotate_vel() 按角度差分档给出"
         "（大于 45° 为 ±0.5 rad/s，大于 15° 为 ±0.4 rad/s，其余 ±0.3 rad/s；角度差在 180° 附近时"
         "固定按 180° 处理，防止转向左右横跳）。")
add_para(doc, "当小车航向与预瞄点航向偏差小于 5°（M_PI_4/9）时判定转弯完成：将该转弯点记录到 "
         "turn_completed_points_（最多保留 3 个，供 /turn_completed 服务查询，避免重复跟踪同一转弯点），"
         "停车一拍后切回 Follow。若 turn_count_ 超过 200（约 10 s 仍未完成转弯），"
         "进入脱困模式：先以 -0.5 m/s 后退 20 个周期，再继续旋转。")
add_heading(doc, "4. Reverse（倒车状态）", level=2)
add_para(doc, "清零两个计数器，以 -0.15 m/s 匀速直线倒车。该状态由紧急避障触发"
         "（avoid_state_ == 4）。")

# ===================== 三、进入各状态的条件 =====================
add_heading(doc, "三、进入状态机各状态的条件")
add_para(doc, "前置安全检查在每个循环周期、状态机 switch 之前执行，满足条件时直接改写 "
         "running_state_ 并跳转至状态机入口，优先级高于 Follow/Turn 内部的正常流转。")

add_heading(doc, "1. 进入 Stop 的条件", level=2)
add_table(doc,
          ["序号", "进入条件", "代码位置 / 说明"],
          [["1", "车辆位置信息丢失：距上次收到 /Mower/position 超过 1 s",
            "state_machine_run() 开头，last_pos_time_ 超时"],
           ["2", "五维路径信息丢失：距上次收到 /lawn_mower/global_path 超过 3 s",
            "last_local_path_time_ 超时"],
           ["3", "紧急预警需停车：避障状态 avoid_state_ == 3",
            "话题 /lawn_mower/avoid_state 更新"],
           ["4", "未找到离车辆最近的路径点（路径为空，findNearestPoint 返回 -1）",
            "最近点查找失败"],
           ["5", "到达路径终点：最近点是已完成转弯的点且为路径最后一个点，停车等待新路径",
            "Follow 状态下的终点判断"]])

add_heading(doc, "2. 进入 Reverse 的条件", level=2)
add_table(doc,
          ["序号", "进入条件", "代码位置 / 说明"],
          [["1", "紧急预警需倒车：避障状态 avoid_state_ == 4",
            "前置安全检查，优先级最高"]])

add_heading(doc, "3. 进入 Turn 的条件", level=2)
add_table(doc,
          ["序号", "进入条件", "代码位置 / 说明"],
          [["1", "Follow 状态下到达预瞄点（预瞄距离 < 0.05 m，或预瞄点 local_x < 0.03 即已越过），"
            "且该预瞄点 gear == 2（转弯点）",
            "Follow 分支内的到达判断"]])

add_heading(doc, "4. 进入 / 回到 Follow 的条件", level=2)
add_table(doc,
          ["序号", "进入条件", "代码位置 / 说明"],
          [["1", "节点启动：构造函数中将 running_state_ 初始化为 Follow", "PurePursuit() 构造函数"],
           ["2", "Stop 状态执行一个周期后自动切回 Follow", "Stop 分支末尾"],
           ["3", "Turn 完成：航向偏差 |local_yaw| < 5°，停车一拍后切回 Follow", "Turn 分支"],
           ["4", "Turn 状态下预瞄点角度刷新失败（原预瞄点已从路径中删除），回退到 Follow",
            "Turn 状态预瞄点角度更新处"]])

# ===================== 四、状态流转 =====================
add_heading(doc, "四、状态流转总览")
add_para(doc,
         "Follow ──到达转弯预瞄点(gear=2)──▶ Turn ──航向偏差<5°──▶ Follow\n"
         "Follow/任意状态 ──位置超时1s / 路径超时3s / avoid_state=3 / 无最近点 / 到达终点──▶ Stop ──下一周期──▶ Follow\n"
         "任意状态 ──avoid_state=4──▶ Reverse", indent=False)
add_para(doc, "说明：代码中 Reverse 状态没有显式出口。避障状态解除后 running_state_ 仍为 Reverse，"
         "每周期继续以 -0.15 m/s 倒车；目前只能依靠 /signal 话题发送 stop / mowing_finished / reset "
         "指令复位（该复位清零各标志位和计数器，但不直接修改 running_state_）。"
         "实际使用中如需退出倒车，需要上层节点配合处理，这一点请留意。")

# ===================== 五、影响控制但不改变状态的条件 =====================
add_heading(doc, "五、不进入状态机但影响控制行为的其他条件")
add_table(doc,
          ["条件", "来源话题", "效果"],
          [["avoid_state_ == 2（较远障碍物普通预警）或 outboundary_flag_ == 1（出边界）",
            "/lawn_mower/avoid_state、/lawn_mower/out_of_bounds",
            "置 low_speed_flag_，期望速度降为低速 low_speed_（默认 0.3 m/s）；转弯点处期望速度固定 0.25 m/s"],
           ["收到停车信号 stop_car_ = true",
            "/mower/stop_car",
            "publishCommand() 发布指令时强制零速度、P 挡"],
           ["收到复位指令 stop / mowing_finished / reset",
            "/signal",
            "清零避障状态、出界标志、路径、计数器、转弯完成点列表及割盘配置"],
           ["割盘控制指令 open / close / 数字 2~11",
            "/signal",
            "设置割盘开关 mover_bool 与割盘高度 mower_height，随控制指令一并下发"]])

doc.save("pure_pursuit状态机说明.docx")
print("已生成 pure_pursuit状态机说明.docx")
